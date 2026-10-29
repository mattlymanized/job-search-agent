#!/usr/bin/env python3
"""Build the three-tab job search tracker workbook.

Usage:
    python3 build_tracker.py --out "/path/to/Job Search Tracker.xlsx"
    python3 build_tracker.py --out tracker.xlsx --rubric rubric.json

rubric.json is optional. It is a list of {"name": str, "weight": number}
objects. Weights are points out of 100 and must sum to 100. Without it, the
default 12-category rubric below is used.

The script refuses to overwrite an existing file unless --force is passed,
because the tracker holds the user's real pipeline.
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DEFAULT_RUBRIC = [
    {"name": "Culture & Leadership", "weight": 12},
    {"name": "Autonomy & Org Influence", "weight": 12},
    {"name": "Compensation", "weight": 12},
    {"name": "Flexibility", "weight": 10},
    {"name": "Management Style", "weight": 8},
    {"name": "Team Structure", "weight": 8},
    {"name": "Impact & Visibility", "weight": 8},
    {"name": "Company Stage & Size", "weight": 8},
    {"name": "Mission & Industry", "weight": 7},
    {"name": "Career Trajectory", "weight": 6},
    {"name": "Product & Buyer", "weight": 5},
    {"name": "Recognition & Feedback", "weight": 4},
]

BASE_COLUMNS = [
    ("Company", 24),
    ("Score", 8),
    ("Score Band", 14),
    ("Status", 20),
    ("Role/Title", 32),
    ("Posting Link", 36),
    ("Date First Seen", 15),
    ("Date Last Confirmed Open", 18),
    ("Criteria Met (Summary)", 44),
    ("Criteria Gaps (Summary)", 44),
    ("Sources", 36),
    ("Notes", 60),
]

STATUSES = [
    "Pursue",
    "Pursue with Diligence",
    "Applied",
    "Followed Up",
    "Call Scheduled",
    "Interview",
    "Pass",
    "Closed",
    "New - Needs Scoring",
    "Deal-Breaker",
    "Declined",
]

SHEETS = ["Tracker", "Passed - Closed", "Applied and Lost"]

HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def load_rubric(path):
    if not path:
        return DEFAULT_RUBRIC
    rubric = json.loads(Path(path).read_text())
    if not isinstance(rubric, list) or not rubric:
        sys.exit("rubric must be a non-empty JSON list")
    for item in rubric:
        if "name" not in item or "weight" not in item:
            sys.exit("every rubric item needs a name and a weight")
    return rubric


def build(out, rubric):
    total = sum(float(c["weight"]) for c in rubric)
    if abs(total - 100) > 0.01:
        sys.exit(f"rubric weights sum to {total}, they must sum to 100")

    wb = Workbook()
    wb.remove(wb.active)

    headers = [(n, w) for n, w in BASE_COLUMNS]
    for c in rubric:
        weight = c["weight"]
        weight = int(weight) if float(weight).is_integer() else weight
        headers.append((f"{c['name']} ({weight}%)", 16))

    last_col = get_column_letter(len(headers))

    for sheet_name in SHEETS:
        ws = wb.create_sheet(sheet_name)
        for idx, (title, width) in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=idx, value=title)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.row_dimensions[1].height = 42
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{last_col}1"

        # Score colors: green 70+, amber 55-69, red below 55.
        ws.conditional_formatting.add(
            "B2:B2000",
            CellIsRule(operator="greaterThanOrEqual", formula=["70"],
                       fill=PatternFill("solid", start_color="C6EFCE")),
        )
        ws.conditional_formatting.add(
            "B2:B2000",
            CellIsRule(operator="between", formula=["55", "69.99"],
                       fill=PatternFill("solid", start_color="FFEB9C")),
        )
        ws.conditional_formatting.add(
            "B2:B2000",
            CellIsRule(operator="lessThan", formula=["55"],
                       fill=PatternFill("solid", start_color="F4CCCC")),
        )

    # Status dropdown on the active tab only.
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(STATUSES) + '"',
        allow_blank=True,
    )
    dv.error = "Pick a status from the list."
    wb["Tracker"].add_data_validation(dv)
    dv.add("D2:D2000")

    wb.save(out)


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", required=True, help="where to write the .xlsx")
    parser.add_argument("--rubric", help="optional rubric JSON file")
    parser.add_argument("--force", action="store_true",
                        help="overwrite an existing file (off by default)")
    args = parser.parse_args()

    out = Path(args.out).expanduser()
    if out.exists() and not args.force:
        sys.exit(f"{out} already exists. Not overwriting. Pass --force only if you mean it.")
    out.parent.mkdir(parents=True, exist_ok=True)

    rubric = load_rubric(args.rubric)
    build(out, rubric)
    print(f"Wrote {out}")
    print(f"Tabs: {', '.join(SHEETS)}")
    print(f"Categories: {len(rubric)} (weights sum to 100)")


if __name__ == "__main__":
    main()
